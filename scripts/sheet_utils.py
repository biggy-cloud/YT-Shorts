#!/usr/bin/env python3
"""CRUD helper for output/content_ideas.xlsx. Used by the 4 subagents so they
never hand-edit the spreadsheet directly (avoids format drift / corruption).

Usage:
  python3 scripts/sheet_utils.py add --link URL --hook "..." --summary "..."
  python3 scripts/sheet_utils.py update --id 3 --field "New Hook Options" --value "..."
  python3 scripts/sheet_utils.py get --id 3
  python3 scripts/sheet_utils.py list
"""
import argparse
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook

SHEET_PATH = Path(__file__).resolve().parent.parent / "output" / "content_ideas.xlsx"

COLUMNS = [
    "ID",
    "Source Link",
    "Original Hook",
    "Original Story Summary",
    "Why It Went Viral",
    "New Hook Options",
    "Chosen Hook",
    "Script",
    "CTA",
    "QA Status",
    "QA Notes",
    "Status",
]


def _ensure_sheet():
    if SHEET_PATH.exists():
        return load_workbook(SHEET_PATH)
    wb = Workbook()
    ws = wb.active
    ws.title = "Ideas"
    ws.append(COLUMNS)
    for i, col in enumerate(COLUMNS, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = max(
            14, min(40, len(col) + 4)
        )
    wb.save(SHEET_PATH)
    return wb


def cmd_add(args):
    wb = _ensure_sheet()
    ws = wb["Ideas"]
    next_id = ws.max_row  # header is row 1, so max_row == count of existing ideas + 1
    row = {
        "ID": next_id,
        "Source Link": args.link or "",
        "Original Hook": args.hook or "",
        "Original Story Summary": args.summary or "",
        "Why It Went Viral": args.why or "",
        "New Hook Options": "",
        "Chosen Hook": "",
        "Script": "",
        "CTA": "",
        "QA Status": "",
        "QA Notes": "",
        "Status": "Researched",
    }
    ws.append([row[c] for c in COLUMNS])
    wb.save(SHEET_PATH)
    print(f"Added idea ID={next_id} -> {SHEET_PATH}")


def _find_row(ws, idea_id):
    for r in range(2, ws.max_row + 1):
        if str(ws.cell(row=r, column=1).value) == str(idea_id):
            return r
    return None


def cmd_update(args):
    wb = _ensure_sheet()
    ws = wb["Ideas"]
    r = _find_row(ws, args.id)
    if r is None:
        print(f"ID {args.id} not found", file=sys.stderr)
        sys.exit(1)
    if args.field not in COLUMNS:
        print(f"Unknown field {args.field!r}. Valid: {COLUMNS}", file=sys.stderr)
        sys.exit(1)
    c = COLUMNS.index(args.field) + 1
    ws.cell(row=r, column=c, value=args.value)
    wb.save(SHEET_PATH)
    print(f"Updated ID={args.id} field={args.field!r}")


def cmd_get(args):
    wb = _ensure_sheet()
    ws = wb["Ideas"]
    r = _find_row(ws, args.id)
    if r is None:
        print(f"ID {args.id} not found", file=sys.stderr)
        sys.exit(1)
    for i, col in enumerate(COLUMNS, start=1):
        print(f"{col}: {ws.cell(row=r, column=i).value}")


def cmd_list(args):
    wb = _ensure_sheet()
    ws = wb["Ideas"]
    if ws.max_row < 2:
        print("(no ideas yet)")
        return
    for r in range(2, ws.max_row + 1):
        idea_id = ws.cell(row=r, column=1).value
        hook = ws.cell(row=r, column=3).value or ""
        status = ws.cell(row=r, column=12).value or ""
        print(f"[{idea_id}] ({status}) {hook[:60]}")


def main():
    p = argparse.ArgumentParser()
    sub = p.add_subparsers(dest="cmd", required=True)

    a = sub.add_parser("add")
    a.add_argument("--link", default="")
    a.add_argument("--hook", default="")
    a.add_argument("--summary", default="")
    a.add_argument("--why", default="")
    a.set_defaults(func=cmd_add)

    u = sub.add_parser("update")
    u.add_argument("--id", required=True)
    u.add_argument("--field", required=True)
    u.add_argument("--value", required=True)
    u.set_defaults(func=cmd_update)

    g = sub.add_parser("get")
    g.add_argument("--id", required=True)
    g.set_defaults(func=cmd_get)

    l = sub.add_parser("list")
    l.set_defaults(func=cmd_list)

    args = p.parse_args()
    args.func(args)


if __name__ == "__main__":
    main()
